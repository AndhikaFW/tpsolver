"""Builds the final Kemjar answer .docx from the official template.

Fills in the identity table, strips the red "Petunjuk"/"Contoh" section that
the template instructs students to remove before submitting, and appends the
user's own answers (plus screenshots and IEEE-style references) formatted to
match the pattern shown in the template's example.
"""

from __future__ import annotations

import copy
import io
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

PART_TITLES = {"1": "Teori", "2": "Praktek"}


def _cell_value_run(table, label: str):
    """Find the value cell's placeholder run immediately following a label cell."""
    for row in table.rows:
        cells = row.cells
        for i, cell in enumerate(cells[:-1]):
            if cell.text.strip() == label:
                value_cell = cells[i + 1]
                for p in value_cell.paragraphs:
                    if p.runs:
                        return p.runs[0]
    return None


def fill_identity(doc: DocumentType, nama: str, npm: str, no_modul: str, tipe: str) -> None:
    identity_table = doc.tables[1]
    for label, value in (
        ("Nama", nama),
        ("NPM", npm),
        ("No. Modul", no_modul),
        ("Tipe", tipe),
    ):
        run = _cell_value_run(identity_table, label)
        if run is not None:
            run.text = value


def _find_style_templates(doc: DocumentType) -> dict[str, Paragraph]:
    templates: dict[str, Paragraph] = {}
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == "Referensi:" and "ref_label" not in templates:
            templates["ref_label"] = p
        elif text.startswith("cisco.com diakses") and "ref_bullet" not in templates:
            templates["ref_bullet"] = p
        elif text.startswith("Lorem Ipsum") and "answer" not in templates:
            templates["answer"] = p
    return templates


def _strip_petunjuk_and_contoh(doc: DocumentType) -> None:
    """Remove the entire red "Petunjuk" block and the "Contoh" example block.

    Only the "Petunjuk:" label and its bullet points are colored red in the
    template; the illustrative Lorem-Ipsum example that follows is not, but
    it is still a placeholder that must not appear in a submitted answer, so
    the whole trailing section (everything after the identity table, up to
    the closing section properties) is removed and replaced with real
    content.
    """
    body = doc.element.body
    start = None
    for i, child in enumerate(body):
        if child.tag == qn("w:p"):
            text = "".join(t.text or "" for t in child.findall(".//" + qn("w:t")))
            if text.strip().startswith("Petunjuk"):
                start = i
                break
    if start is None:
        return
    to_remove = []
    for child in list(body)[start:]:
        if child.tag == qn("w:sectPr"):
            break
        to_remove.append(child)
    for el in to_remove:
        body.remove(el)


def _clone_pPr(template_paragraph: Paragraph, strip_numPr: bool = False):
    pPr = template_paragraph._p.find(qn("w:pPr"))
    if pPr is None:
        return None
    new_pPr = copy.deepcopy(pPr)
    if strip_numPr:
        numPr = new_pPr.find(qn("w:numPr"))
        if numPr is not None:
            new_pPr.remove(numPr)
    return new_pPr


def _clone_rPr(template_run):
    rPr = template_run._r.find(qn("w:rPr"))
    return copy.deepcopy(rPr) if rPr is not None else None


def _add_styled_paragraph(
    doc: DocumentType, template_paragraph: Paragraph, text: str, strip_numPr: bool = False
) -> Paragraph:
    p = doc.add_paragraph()
    new_pPr = _clone_pPr(template_paragraph, strip_numPr=strip_numPr)
    if new_pPr is not None:
        old = p._p.find(qn("w:pPr"))
        if old is not None:
            p._p.remove(old)
        p._p.insert(0, new_pPr)
    run = p.add_run(text)
    if template_paragraph.runs:
        rPr = _clone_rPr(template_paragraph.runs[0])
        if rPr is not None:
            old_rpr = run._r.find(qn("w:rPr"))
            if old_rpr is not None:
                run._r.remove(old_rpr)
            run._r.insert(0, rPr)
    return p


def _add_heading(doc: DocumentType, ref_label_template: Paragraph, text: str) -> Paragraph:
    """A bold paragraph in the same 12pt Times New Roman as the rest of the
    document (per the template's font rule) but flush with the left margin,
    used to mark each "Part N" section since the template has no such
    heading of its own."""
    p = _add_styled_paragraph(doc, ref_label_template, text)
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None:
        ind = pPr.find(qn("w:ind"))
        if ind is not None:
            pPr.remove(ind)
    return p


def _add_image(doc: DocumentType, image_bytes: bytes) -> None:
    p = doc.add_paragraph()
    p.alignment = 1  # center
    run = p.add_run()
    run.add_picture(io.BytesIO(image_bytes), width=Inches(6.0))


def build_output(
    template_bytes: bytes,
    identity: dict[str, str],
    parts_answers: dict[str, list[dict[str, Any]]],
    images_by_item: dict[str, list[bytes]],
) -> bytes:
    """Assemble the final .docx.

    parts_answers: {"1": [{"number": 1, "answer": str, "references": [str, ...]}, ...], ...}
    images_by_item: {"2-1": [bytes, ...], ...} keyed "{part}-{number}" for Part 2 screenshots.
    """
    doc = Document(io.BytesIO(template_bytes))

    fill_identity(
        doc,
        identity.get("nama", ""),
        identity.get("npm", ""),
        identity.get("no_modul", ""),
        identity.get("tipe", "TP"),
    )

    templates = _find_style_templates(doc)
    _strip_petunjuk_and_contoh(doc)

    answer_tmpl = templates["answer"]
    ref_label_tmpl = templates["ref_label"]
    ref_bullet_tmpl = templates["ref_bullet"]

    for part_number in ("1", "2"):
        items = parts_answers.get(part_number) or []
        if not items:
            continue
        _add_heading(doc, ref_label_tmpl, f"Part {part_number} - {PART_TITLES[part_number]}")

        for item in items:
            number = item["number"]
            answer_text = (item.get("answer") or "").strip()
            lines = answer_text.split("\n") if answer_text else [""]

            _add_styled_paragraph(doc, answer_tmpl, f"{number}. {lines[0]}", strip_numPr=True)
            for line in lines[1:]:
                if line.strip():
                    _add_styled_paragraph(doc, answer_tmpl, line, strip_numPr=True)

            for image_bytes in images_by_item.get(f"{part_number}-{number}", []):
                _add_image(doc, image_bytes)

            references = item.get("references") or []
            if part_number == "1" and references:
                _add_styled_paragraph(doc, ref_label_tmpl, "Referensi:")
                for ref in references:
                    if ref.strip():
                        _add_styled_paragraph(doc, ref_bullet_tmpl, ref.strip())

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
